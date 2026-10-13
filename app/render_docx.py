"""Render structured document content as an editable Word file."""
from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml.ns import qn

from .models import DocIR, Source, cited_ids


def _style(doc, name: str, fallback: str = "Normal") -> str:
    try:
        doc.styles[name]
        return name
    except KeyError:
        return fallback


def render_docx(ir: DocIR, sources: list[Source], template: Optional[Path], out: Path) -> Path:
    doc = Document(str(template)) if template else Document()
    body = doc.element.body
    for el in list(body):  # clear template body, keep section properties (page size, margins, header/footer)
        if el.tag != qn("w:sectPr"):
            body.remove(el)

    doc.add_paragraph(ir.title, style=_style(doc, "Title"))
    if ir.subtitle:
        doc.add_paragraph(ir.subtitle, style=_style(doc, "Subtitle"))

    for s in ir.sections:
        doc.add_paragraph(s.heading, style=_style(doc, f"Heading {max(1, min(s.level, 3))}"))
        for p in s.paragraphs:
            doc.add_paragraph(p, style="Normal")
        for b in s.bullets:
            doc.add_paragraph(b, style=_style(doc, "List Bullet"))

    ids = cited_ids(doc=ir)
    used = [s for s in sources if s.id in ids]
    if used:
        doc.add_paragraph("References", style=_style(doc, "Heading 1"))
        for s in used:
            doc.add_paragraph(f"[{s.id}] {s.title}" + (f" — {s.url}" if s.url else f" ({s.kind})"), style="Normal")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out

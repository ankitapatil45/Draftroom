"""Render structured slide content as an editable PowerPoint file."""
from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.util import Inches, Pt

from .models import DeckIR, Source, cited_ids

BODY_TYPES = {PH.BODY, PH.OBJECT}


def _types(layout):
    return [p.placeholder_format.type for p in layout.placeholders]


def pick_layouts(prs):
    title = content = None
    for l in prs.slide_layouts:
        t = _types(l)
        if title is None and (PH.CENTER_TITLE in t or (PH.TITLE in t and PH.SUBTITLE in t)):
            title = l
        if content is None and PH.TITLE in t and any(x in BODY_TYPES for x in t):
            content = l
    layouts = list(prs.slide_layouts)
    return title or layouts[0], content or (layouts[1] if len(layouts) > 1 else layouts[0])


def _fill_body(slide, bullets: list[str]):
    body = next((p for p in slide.placeholders if p.placeholder_format.type in BODY_TYPES), None)
    if body is None:
        w, h = slide.part.package.presentation_part.presentation.slide_width, \
               slide.part.package.presentation_part.presentation.slide_height
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), w - Inches(1.4), h - Inches(2.2))
    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    total = sum(len(b) for b in bullets)
    size = None
    if len(bullets) > 6 or total > 450:
        size = Pt(16)
    elif len(bullets) > 4 or total > 320:
        size = Pt(18)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        if size:
            for r in p.runs:
                r.font.size = size


def render_pptx(ir: DeckIR, sources: list[Source], template: Optional[Path], out: Path) -> Path:
    prs = Presentation(str(template)) if template else Presentation()
    lst = prs.slides._sldIdLst
    for sid in list(lst):  # drop the template's sample slides, keep master/layouts/theme
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    title_l, content_l = pick_layouts(prs)

    s = prs.slides.add_slide(title_l)
    if s.shapes.title is not None:
        s.shapes.title.text = ir.title
    sub = next((p for p in s.placeholders if p.placeholder_format.type == PH.SUBTITLE), None)
    if sub is not None:
        sub.text = ir.subtitle

    slides = [x for x in ir.slides if x.kind != "title"]
    ids = cited_ids(deck=ir)
    used = [x for x in sources if x.id in ids]
    for sl in slides:
        if sl.kind == "sources":
            continue
        s = prs.slides.add_slide(content_l)
        if s.shapes.title is not None:
            s.shapes.title.text = sl.title
        _fill_body(s, sl.bullets)
        if sl.notes:
            s.notes_slide.notes_text_frame.text = sl.notes
    if used:
        s = prs.slides.add_slide(content_l)
        if s.shapes.title is not None:
            s.shapes.title.text = "Sources"
        _fill_body(s, [f"[{x.id}] {x.title}" + (f" — {x.url}" if x.url else "") for x in used][:8])
        for p in s.placeholders:
            if p.placeholder_format.type in BODY_TYPES:
                for para in p.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size = Pt(12)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out

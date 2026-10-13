"""Extract text and style details from supported reference files."""
import io
import re
import zipfile
from collections import Counter
from pathlib import Path

from .llm import llm

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def kind_of(name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        return "docx"
    if ext == ".pdf":
        return "pdf"
    if ext == ".pptx":
        return "pptx"
    if ext in IMAGE_EXT:
        return "image"
    return "text"


# ------------------------------------------------------------------ OCR / vision
def ocr_image(img) -> str:
    try:
        import pytesseract

        return pytesseract.image_to_string(img).strip()
    except Exception:
        return ""


def _image_text(path: Path) -> str:
    from PIL import Image

    img = Image.open(path)
    text = ocr_image(img)
    if len(text) < 40 and llm.live:  # OCR weak/unavailable -> vision model
        buf = io.BytesIO()
        img.convert("RGB").save(buf, "PNG")
        text = llm.describe_image(buf.getvalue(), "image/png") or text
    return text


def _pdf_text(path: Path) -> str:
    import pdfplumber

    pages = []
    with pdfplumber.open(path) as pdf:
        for p in pdf.pages:
            pages.append((p.extract_text() or "").strip())
    if sum(len(p) for p in pages) < 50:  # scanned PDF -> render pages and OCR / vision
        import pypdfium2 as pdfium

        pages = []
        pdf = pdfium.PdfDocument(str(path))
        for i in range(min(len(pdf), 30)):
            img = pdf[i].render(scale=2).to_pil()
            t = ocr_image(img)
            if len(t) < 40 and llm.live:
                buf = io.BytesIO()
                img.convert("RGB").save(buf, "PNG")
                t = llm.describe_image(buf.getvalue(), "image/png")
            pages.append(t)
    return "\n\n".join(pages)


def _docx_text(path: Path) -> str:
    from docx import Document

    d = Document(str(path))
    out = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            out.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(out)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    out = []
    for i, s in enumerate(Presentation(str(path)).slides, 1):
        out.append(f"[Slide {i}]")
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
    return "\n".join(out)


def extract_text(path: str | Path) -> str:
    path = Path(path)
    k = kind_of(path.name)
    if k == "docx":
        return _docx_text(path)
    if k == "pdf":
        return _pdf_text(path)
    if k == "pptx":
        return _pptx_text(path)
    if k == "image":
        return _image_text(path)
    text = path.read_text(errors="ignore")
    return re.sub(r"^#+\s*(.+?)\s*$", r"\1.", text, flags=re.M)  # markdown headings -> plain sentences


# ------------------------------------------------------------------ tone
def tone_profile(text: str) -> dict:
    sents = [s for s in re.split(r"(?<=[.!?])\s+", text) if len(s.split()) > 2]
    words = re.findall(r"[A-Za-z']+", text.lower())
    n = max(len(words), 1)
    we = sum(words.count(w) for w in ("we", "our", "us"))
    you = sum(words.count(w) for w in ("you", "your"))
    contr = sum(1 for w in words if "'" in w)
    avg = sum(len(s.split()) for s in sents) / max(len(sents), 1)
    prof = {
        "avg_sentence_words": round(avg, 1),
        "voice": "first-person plural (we/our)" if we > you else ("second person (you)" if you > we else "neutral/third person"),
        "register": "formal" if contr / n < 0.005 else "conversational",
        "sentence_length": "long, detailed" if avg > 22 else ("medium" if avg > 14 else "short, punchy"),
    }
    sample = " ".join(sents[:4])[:500]
    if llm.live and sample:
        d = llm.complete("You are an editor.", "Describe the tone, voice and writing style of this text in 2 sentences "
                         "so another writer could imitate it:\n" + text[:3000], 300)
        prof["description"] = d.strip()
    else:
        prof["description"] = (f"{prof['register'].capitalize()} register, {prof['voice']}, "
                               f"{prof['sentence_length']} sentences (~{prof['avg_sentence_words']} words).")
    prof["sample"] = sample
    return prof


# ------------------------------------------------------------------ DOCX style analysis
def analyze_docx(path: str | Path) -> dict:
    from docx import Document

    d = Document(str(path))
    sec = d.sections[0]
    styles = Counter(p.style.name for p in d.paragraphs if p.text.strip())
    headings = []
    for p in d.paragraphs:
        m = re.match(r"Heading (\d)", p.style.name)
        if m and p.text.strip():
            headings.append({"level": int(m.group(1)), "text": p.text.strip()})

    def font(style_name):
        try:
            st = d.styles[style_name]
            f = st.font
            return {"name": f.name, "size": f.size.pt if f.size else None,
                    "bold": f.bold, "color": str(f.color.rgb) if f.color and f.color.type and f.color.rgb else None}
        except KeyError:
            return None

    body = [p.text for p in d.paragraphs if p.style.name in ("Normal", "Body Text") and len(p.text.split()) > 8]
    text = "\n".join(body)
    return {
        "kind": "docx",
        "page": {"width_in": round(sec.page_width.inches, 2) if sec.page_width else None,
                 "height_in": round(sec.page_height.inches, 2) if sec.page_height else None,
                 "margins_in": [round(x.inches, 2) if x is not None else None
                                for x in (sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin)]},
        "paragraph_styles_used": dict(styles),
        "fonts": {n: font(n) for n in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet")},
        "headings": headings,
        "tables": len(d.tables),
        "images": len(d.inline_shapes),
        "header_text": " ".join(p.text for p in sec.header.paragraphs).strip(),
        "footer_text": " ".join(p.text for p in sec.footer.paragraphs).strip(),
        "words": len(_docx_text(Path(path)).split()),
        "tone": tone_profile(text or _docx_text(Path(path))),
    }


# ------------------------------------------------------------------ PPTX style analysis
def analyze_pptx(path: str | Path) -> dict:
    from pptx import Presentation

    prs = Presentation(str(path))
    layouts = [{"name": l.name, "placeholders": [p.placeholder_format.type.name if p.placeholder_format.type else "?"
                                                  for p in l.placeholders]} for l in prs.slide_layouts]
    slides, texts = [], []
    for i, s in enumerate(prs.slides, 1):
        title = s.shapes.title.text if s.shapes.title is not None else ""
        bullets = 0
        pics = 0
        for sh in s.shapes:
            if sh.shape_type == 13:
                pics += 1
            if sh.has_text_frame and sh != s.shapes.title:
                bullets += len([p for p in sh.text_frame.paragraphs if p.text.strip()])
                texts.append(sh.text_frame.text)
        slides.append({"n": i, "layout": s.slide_layout.name, "title": title, "bullets": bullets, "pictures": pics})
    xml = ""
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.startswith("ppt/theme/theme") and n.endswith(".xml"):
                xml = z.read(n).decode("utf8", "ignore")
                break
    fonts = re.findall(r'<a:latin typeface="([^"]*)"', xml)[:2]
    colors = re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)[:8]
    avg_b = sum(s["bullets"] for s in slides) / max(len(slides), 1)
    return {
        "kind": "pptx",
        "slide_size_in": [round(prs.slide_width.inches, 2), round(prs.slide_height.inches, 2)],
        "layouts": layouts,
        "slides": slides,
        "theme_fonts": {"major": fonts[0] if fonts else None, "minor": fonts[1] if len(fonts) > 1 else None},
        "theme_colors": colors,
        "avg_bullets_per_slide": round(avg_b, 1),
        "tone": tone_profile("\n".join(texts) or "No body text."),
    }


def analyze_any(path: str | Path) -> dict:
    k = kind_of(Path(path).name)
    if k == "docx":
        return analyze_docx(path)
    if k == "pptx":
        return analyze_pptx(path)
    text = extract_text(path)
    return {"kind": k, "words": len(text.split()), "tone": tone_profile(text),
            "headings": [{"level": 1, "text": l.strip("# ").strip()} for l in text.splitlines()
                         if l.strip() and len(l.split()) < 8 and l.istitle()][:12]}

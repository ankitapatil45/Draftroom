"""End-to-end test, fully offline: no LLM key, web search stubbed, local vector store."""
import sys
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

FAKE_WEB = [
    {"title": "GenAI trends 2026", "url": "https://example.com/a",
     "snippet": "Agentic workflows are the leading enterprise GenAI trend. Competitive vendors differentiate on governance. "
                "Spend on retrieval grounded assistants keeps growing."},
    {"title": "Enterprise AI market", "url": "https://example.com/b",
     "snippet": "Multimodal models reduce document processing time. Market leaders bundle platform credits with models. "
                "Cost control remains a top concern."},
]


@pytest.fixture()
def env(tmp_path, monkeypatch):
    from app.config import settings
    settings.data_dir = tmp_path / "data"
    settings.pinecone_api_key = ""
    settings.openrouter_api_key = ""
    from app import rag, web
    rag.reset_kb()
    monkeypatch.setattr(web, "search", lambda q, n=6: [dict(x) for x in FAKE_WEB])
    monkeypatch.setattr(web, "fetch_text", lambda u, limit=0: "")
    kb = rag.get_kb()
    for p in (ROOT / "examples" / "knowledge_base").iterdir():
        kb.ingest_file(p)
    from app.agents.supervisor import Supervisor
    from app.store import SessionStore
    store = SessionStore()
    sup = Supervisor(store)
    sid = store.create()
    for p in (ROOT / "examples" / "templates").iterdir():
        from app.parsers import kind_of
        store.save_upload(sid, p.name, p.read_bytes(), kind_of(p.name))
    return store, sup, sid


def _open(store, sid, n):
    d = store.version_dir(sid, n)
    return Document(str(d / "document.docx")), Presentation(str(d / "presentation.pptx"))


def test_create_edit_convert_versions(env):
    store, sup, sid = env
    r = sup.handle(sid, "Research the latest Generative AI trends and create a proposal and 12-slide presentation "
                        "using the same tone and style as the uploaded files.")
    assert r["version"] == 1 and set(r["files"]) == {"docx", "pptx"}
    doc, prs = _open(store, sid, 1)
    assert len(prs.slides) == 12, len(prs.slides)
    # template style preserved: custom heading font and header text from the uploaded proposal
    assert doc.styles["Heading 1"].font.name == "Georgia"
    assert "Acme Consulting" in doc.sections[0].header.paragraphs[0].text
    assert prs.slide_width == Presentation(str(Path(store.dir(sid)) / "uploads" / "Company_Template.pptx")).slide_width
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[1]" in text and "References" in text  # citations + reference list
    assert r["validation"]["passed"], r["validation"]

    r2 = sup.handle(sid, "Add an executive summary.")
    assert r2["version"] == 2
    doc2, prs2 = _open(store, sid, 2)
    assert doc2.paragraphs and any(p.text == "Executive Summary" for p in doc2.paragraphs)
    assert doc2.styles["Heading 1"].font.name == "Georgia"

    r3 = sup.handle(sid, "Add a competitive analysis section.")
    assert r3["version"] == 3
    assert any(p.text == "Competitive Analysis" for p in _open(store, sid, 3)[0].paragraphs)

    r4 = sup.handle(sid, "Make the presentation more concise.")
    assert r4["version"] == 4
    for s in _open(store, sid, 4)[1].slides:
        body = [sh for sh in s.placeholders if sh.placeholder_format.idx == 1 and sh.has_text_frame]
        for b in body:
            assert len([p for p in b.text_frame.paragraphs if p.text.strip()]) <= 8

    r5 = sup.handle(sid, "Update the report using the latest web information.")
    assert r5["version"] == 5

    vs = store.versions(sid)
    assert [v["number"] for v in vs] == [1, 2, 3, 4, 5] and vs[1]["parent"] == 1
    assert store.get_trace(sid)  # traceability: every agent step logged
    agents = {t["agent"] for t in store.get_trace(sid)}
    assert {"supervisor", "template_analyst", "web_research", "rag", "doc_writer", "deck_writer", "validator",
            "editor", "renderer"} <= agents

    r6 = sup.handle(sid, "Convert the document to a presentation.")
    assert r6["version"] == 6
    store.revert(sid, 1)
    assert store.latest(sid)["number"] == 7


def test_rag_retrieval(env):
    from app.rag import get_kb
    hits = get_kb().search("vendor neutral competitive advantage", 3)
    assert hits and "competitive" in hits[0]["title"].lower()
